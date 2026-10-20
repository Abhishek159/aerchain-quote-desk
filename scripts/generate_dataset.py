"""Generate the fabricated RFx dataset.

Everything a vendor "sends" is rendered from one hidden truth table, so the
extraction pipeline can be scored exactly against data/eval/ground_truth.json.
The app itself never reads the eval folder.

Run:  uv run python scripts/generate_dataset.py
"""
import csv
import json
import math
import random
import shutil
from datetime import date
from pathlib import Path

import numpy as np
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from PIL import Image, ImageDraw, ImageFilter, ImageFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
rng = random.Random(14)

FX_USD_INR = 88.0  # the buyer's assumed rate; the app shows it and lets you change it
FONT_DIR = Path("/System/Library/Fonts/Supplemental")

BUYER = {
    "name": "Meghna Foods Pvt. Ltd.",
    "plant": "Chakan Plant, Plot B-14, MIDC Phase II, Chakan, Pune 410501",
    "contact": "Ritika Deshmukh, Category Manager - Packaging",
    "email": "ritika.deshmukh@meghnafoods.example",
}
RFX_ID = "MFPL/PKG/RFQ/2026-27/014"

# --------------------------------------------------------------------------
# 1. Line items
# --------------------------------------------------------------------------
BOARD = {
    3: {"spec": "Kraft 150 / Fluting 120 (B) / Kraft 150 GSM, BF 18", "gsm": 150 + 120 * 1.4 + 150},
    5: {"spec": "Kraft 180 / Fl 120 (B) / Kraft 150 / Fl 120 (C) / Kraft 180 GSM, BF 20", "gsm": 180 + 168 + 150 + 168 + 180},
    7: {"spec": "Kraft 200 / Fl 150 / 180 / Fl 150 / 180 / Fl 150 / Kraft 200 GSM, BF 24", "gsm": 200 + 210 + 180 + 210 + 180 + 210 + 200},
}
KG_RATE = {3: 46.0, 5: 48.0, 7: 50.0}
PRINT_ADDER = {"Plain": 0.0, "1 colour": 0.45, "2 colour": 0.9, "4 colour": 3.2}

# id, description, L, W, H, ply, print, annual qty, kind
RAW_LINES = [
    ("L01", "Shipper carton - Namkeen 40g x 24", 200, 150, 100, 3, "Plain", 240000, "RSC"),
    ("L02", "Shipper carton - Namkeen 80g x 24", 250, 180, 120, 3, "1 colour", 180000, "RSC"),
    ("L03", "Shipper carton - Chips 52g x 30", 300, 200, 150, 3, "1 colour", 150000, "RSC"),
    ("L04", "Shipper carton - Chips family pack x 12", 320, 240, 180, 3, "2 colour", 120000, "RSC"),
    ("L05", "Shipper carton - Rusk 300g x 16", 350, 250, 200, 3, "1 colour", 100000, "RSC"),
    ("L06", "Shipper carton - Mixture 1kg x 10", 380, 280, 220, 3, "Plain", 90000, "RSC"),
    ("L07", "Shipper carton - Cookies tin x 12", 400, 300, 200, 3, "2 colour", 80000, "RSC"),
    ("L08", "Shipper carton - Assorted gift x 8", 420, 320, 250, 3, "1 colour", 60000, "RSC"),
    ("L09", "Shipper carton - Bulk namkeen 5kg x 4", 450, 350, 250, 3, "Plain", 50000, "RSC"),
    ("L10", "Shipper carton - Institutional pack", 480, 360, 300, 3, "Plain", 40000, "RSC"),
    ("L11", "Master carton - Biscuit 75g x 48", 300, 250, 200, 5, "1 colour", 150000, "RSC"),
    ("L12", "Master carton - Biscuit 150g x 36", 350, 280, 220, 5, "2 colour", 140000, "RSC"),
    ("L13", "Master carton - Cream biscuit 120g x 40", 380, 300, 250, 5, "1 colour", 120000, "RSC"),
    ("L14", "Master carton - Glucose 800g x 12", 400, 300, 300, 5, "1 colour", 110000, "RSC"),
    ("L15", "Master carton - Juice 200ml x 27", 420, 320, 280, 5, "2 colour", 100000, "RSC"),
    ("L16", "Master carton - Juice 1L x 12", 450, 350, 300, 5, "1 colour", 90000, "RSC"),
    ("L17", "Master carton - Soda 600ml x 24", 480, 360, 320, 5, "1 colour", 80000, "RSC"),
    ("L18", "Master carton - Water 1L x 12", 500, 380, 320, 5, "Plain", 70000, "RSC"),
    ("L19", "Gift pack shipper - Diwali hamper", 520, 400, 350, 5, "4 colour", 50000, "RSC"),
    ("L20", "Master carton - Ghee 1L x 12", 550, 420, 380, 5, "1 colour", 40000, "RSC"),
    ("L21", "Master carton - Atta 5kg x 4", 600, 400, 400, 5, "Plain", 30000, "RSC"),
    ("L22", "Master carton - Sugar 5kg x 5", 600, 450, 450, 5, "1 colour", 25000, "RSC"),
    ("L23", "Heavy-duty carton - Edible oil 15L tin", 400, 400, 350, 7, "1 colour", 30000, "RSC"),
    ("L24", "Heavy-duty carton - Oil 5L jar x 4", 450, 450, 400, 7, "1 colour", 25000, "RSC"),
    ("L25", "Heavy-duty carton - Bulk ingredients 20kg", 600, 500, 500, 7, "Plain", 15000, "RSC"),
    ("L26", "Export carton - Pallet master (new SKU)", 800, 600, 600, 7, "Plain", 8000, "RSC"),
    ("L27", "Partition set 12-cell (fits L11)", 0, 0, 0, 3, "Plain", 200000, "PARTITION12"),
    ("L28", "Partition set 24-cell (fits L16)", 0, 0, 0, 3, "Plain", 120000, "PARTITION24"),
    ("L29", "Layer pad 590 x 390 mm", 590, 390, 0, 3, "Plain", 150000, "PAD"),
    ("L30", "Die-cut tray 400 x 300 x 80 (new SKU)", 400, 300, 80, 3, "2 colour", 100000, "TRAY"),
]


def area_m2(kind, L, W, H):
    if kind == "RSC":
        return (2 * L + 2 * W + 35) * (H + W) / 1e6
    if kind == "PARTITION12":
        return 0.18
    if kind == "PARTITION24":
        return 0.32
    if kind == "PAD":
        return L * W / 1e6
    if kind == "TRAY":
        return (L + 2 * H) * (W + 2 * H) / 1e6
    raise ValueError(kind)


LINES = []
for lid, desc, L, W, H, ply, prn, qty, kind in RAW_LINES:
    wt = area_m2(kind, L, W, H) * BOARD[ply]["gsm"] / 1000
    base = wt * KG_RATE[ply] + 0.5 + PRINT_ADDER[prn]
    if kind.startswith("PARTITION"):
        base += 1.2  # slotting labour
    if kind == "TRAY":
        base += 1.5  # die-cutting
    dims = f"{L} x {W} x {H}" if kind in ("RSC", "TRAY") else (f"{L} x {W}" if kind == "PAD" else "as per drawing")
    LINES.append({
        "line_id": lid,
        "item_code": f"MF-PK-{int(lid[1:]):04d}",
        "description": desc,
        "kind": kind,
        "dimensions_mm": dims,
        "L": L, "W": W, "H": H,
        "ply": ply,
        "board_spec": BOARD[ply]["spec"],
        "print": prn,
        "uom": "piece" if not kind.startswith("PARTITION") else "set",
        "annual_qty": int(round(qty * 0.63 / 1000) * 1000),
        "nominal_weight_kg": round(wt, 3),
        "_base": base,
    })
LINE = {l["line_id"]: l for l in LINES}

# --------------------------------------------------------------------------
# 2. Hidden truth: what each vendor really charges (INR / piece, ex-GST,
#    before discounts, before freight).  Rendered into each vendor's format.
# --------------------------------------------------------------------------
VENDORS = {
    "V1": {"name": "Pacifica Packaging India Pvt. Ltd.", "short": "Pacifica", "factor": 1.03,
           "email": "bids@pacificapack.example", "contact": "Daniel Tan, Regional Sales"},
    "V2": {"name": "Shree Ganesh Packaging Industries", "short": "Shree Ganesh", "factor": 0.985,
           "email": "sales@shreeganeshpack.example", "contact": "Prakash Kulkarni, Proprietor"},
    "V3": {"name": "Konkan Kraft Boxes LLP", "short": "Konkan Kraft", "factor": 1.015,
           "email": "commercial@konkankraft.example", "contact": "Anjali Sawant, Commercial Head"},
    "V4": {"name": "Pioneer Packwell", "short": "Pioneer", "factor": 0.915,
           "email": "pioneerpackwell@mail.example", "contact": "Salim Shaikh"},
    "V5": {"name": "Maruti Corrupack Pvt. Ltd.", "short": "Maruti", "factor": None,
           "email": "rajesh@maruticorrupack.example", "contact": "Rajesh Agarwal, Director"},
}


def vendor_price(vid, lid):
    f = VENDORS[vid]["factor"]
    return LINE[lid]["_base"] * f * (1 + rng.gauss(0, 0.045))


# Last year's prices (FY 2025-26).  Maruti was the incumbent on almost everything.
LAST_YEAR = {}
for l in LINES:
    if l["line_id"] in ("L26", "L30"):
        continue  # new SKUs, never bought before
    LAST_YEAR[l["line_id"]] = round(l["_base"] * 0.965 * (1 + rng.gauss(0, 0.03)), 2)

truth = {"fx_usd_inr": FX_USD_INR, "vendors": {}}


def T(vid):
    return truth["vendors"].setdefault(vid, {"lines": {}, "terms": {}, "questionnaire": {}})


def record(vid, lid, status, native, native_unit, currency, inr_per_unit_exgst, note=""):
    T(vid)["lines"][lid] = {
        "status": status, "native_price": native, "native_unit": native_unit, "currency": currency,
        "inr_per_unit_ex_gst": None if inr_per_unit_exgst is None else round(inr_per_unit_exgst, 4),
        "note": note,
    }


# --------------------------------------------------------------------------
# 3. Questionnaire
# --------------------------------------------------------------------------
QUESTIONNAIRE = [
    {"qid": "Q1", "rule": {"op": "yes_and_valid_cert", "cert": "ISO 9001"}, "text": "Do you hold a valid ISO 9001 certificate? Please attach a copy.", "type": "yes_no_attachment",
     "criterion": "Must be Yes, with a certificate valid on the award date", "knockout": True},
    {"qid": "Q2", "rule": {"op": "yes"}, "text": "Do you test every lot for BCT / ECT / burst strength (in-house lab or NABL-accredited lab)?", "type": "yes_no",
     "criterion": "Must be Yes", "knockout": True},
    {"qid": "Q3", "rule": {"op": ">=", "value": 150}, "text": "Installed corrugation capacity (MT per month)", "type": "number",
     "criterion": ">= 150 MT/month", "knockout": True},
    {"qid": "Q4", "rule": {"op": "<=", "value": 1.5}, "text": "Customer rejection rate (% of supplied quantity) over the last 12 months", "type": "number",
     "criterion": "<= 1.5%", "knockout": True},
    {"qid": "Q5", "text": "Standard lead time from PO to delivery at Chakan (days)", "type": "number",
     "criterion": "<= 21 days preferred", "knockout": False},
    {"qid": "Q6", "text": "Are printing inks low-migration / compliant with IS 15495 for food packaging?", "type": "yes_no",
     "criterion": "Yes preferred", "knockout": False},
    {"qid": "Q7", "text": "Recycled fibre content of the board (%)", "type": "number",
     "criterion": "Information only", "knockout": False},
    {"qid": "Q8", "text": "Number of FMCG / food customers currently supplied", "type": "number",
     "criterion": "Information only", "knockout": False},
    {"qid": "Q9", "text": "Distance of manufacturing plant from Chakan (km)", "type": "number",
     "criterion": "Information only", "knockout": False},
    {"qid": "Q10", "text": "Do you accept 60-day payment terms?", "type": "yes_no",
     "criterion": "Yes preferred", "knockout": False},
]

TERMS_REQUESTED = {
    "currency": "INR",
    "price_basis": "Per piece (or per set for partitions), ex-GST, delivered to Chakan plant",
    "payment_terms": "60 days from receipt of material",
    "validity": "Prices firm for 12 months (Nov 2026 - Oct 2027)",
    "delivery": "Call-off against monthly schedule, 14-21 day lead time",
    "quote_due": "2026-10-10",
    "gst": "Show GST separately",
}

# --------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------

def fresh(path: Path):
    if path.exists():
        shutil.rmtree(path)
    path.mkdir(parents=True)
    return path


def write_email(folder, vid, subject, sent, body, attachments):
    meta = {"from_name": VENDORS[vid]["contact"], "from_company": VENDORS[vid]["name"],
            "from_email": VENDORS[vid]["email"], "to": BUYER["email"], "subject": subject,
            "received_at": sent, "body": body, "attachments": attachments}
    (folder / "email.json").write_text(json.dumps(meta, indent=2, ensure_ascii=False))


def font(name, size):
    return ImageFont.truetype(str(FONT_DIR / name), size)


def iso_certificate(path, company, cert_no, issued, expires, scope):
    c = canvas.Canvas(str(path), pagesize=A4)
    w, h = A4
    c.setStrokeColor(colors.HexColor("#1f3b6f"))
    c.setLineWidth(4)
    c.rect(15 * mm, 15 * mm, w - 30 * mm, h - 30 * mm)
    c.setFont("Helvetica-Bold", 13)
    c.drawCentredString(w / 2, h - 40 * mm, "QUALITAS CERTIFICATION SERVICES (fictional registrar)")
    c.setFont("Helvetica-Bold", 30)
    c.drawCentredString(w / 2, h - 70 * mm, "CERTIFICATE")
    c.setFont("Helvetica", 12)
    c.drawCentredString(w / 2, h - 85 * mm, "This is to certify that the Quality Management System of")
    c.setFont("Helvetica-Bold", 18)
    c.drawCentredString(w / 2, h - 100 * mm, company)
    c.setFont("Helvetica", 12)
    c.drawCentredString(w / 2, h - 115 * mm, "has been assessed and found to conform to the requirements of")
    c.setFont("Helvetica-Bold", 20)
    c.drawCentredString(w / 2, h - 132 * mm, "ISO 9001:2015")
    c.setFont("Helvetica", 11)
    c.drawCentredString(w / 2, h - 148 * mm, f"Scope: {scope}")
    y = h - 180 * mm
    for k, v in [("Certificate No.", cert_no), ("Date of issue", issued), ("Valid until", expires)]:
        c.drawString(45 * mm, y, k)
        c.drawString(100 * mm, y, ": " + v)
        y -= 9 * mm
    c.setFont("Helvetica-Oblique", 9)
    c.drawCentredString(w / 2, 30 * mm, "Validity of this certificate is subject to annual surveillance audits.")
    c.save()


# --------------------------------------------------------------------------
# 4. The RFx itself (what the buyer sent) + quote template
# --------------------------------------------------------------------------

def build_rfx():
    out = fresh(DATA / "rfx")
    rfx = {
        "rfx_id": RFX_ID, "title": "Corrugated boxes & accessories - annual rate contract FY 2026-27",
        "buyer": BUYER, "issued_on": "2026-10-01", "due_on": "2026-10-10",
        "category": "Corrugated packaging",
        "scope": "Annual rate contract for 30 corrugated packaging SKUs (3/5/7-ply RSC cartons, partitions, "
                 "layer pads, die-cut trays) for the Chakan plant. Quantities are annual estimates; "
                 "call-offs against monthly schedules.",
        "lines": [{k: v for k, v in l.items() if not k.startswith("_") and k not in ("L", "W", "H")} for l in LINES],
        "questionnaire": QUESTIONNAIRE,
        "terms": TERMS_REQUESTED,
        "vendors_invited": [{"vendor_id": k, "name": v["name"], "email": v["email"]} for k, v in VENDORS.items()],
    }
    (out / "rfx.json").write_text(json.dumps(rfx, indent=2, ensure_ascii=False))

    wb = Workbook()
    ws = wb.active
    ws.title = "Price Bid"
    ws.append([f"{BUYER['name']} - {RFX_ID}"])
    ws.append(["Please quote in INR per piece, ex-GST, delivered Chakan. Do not change the layout."])
    ws.append([])
    hdr = ["Line", "Item code", "Description", "Size (mm)", "Ply", "Board spec", "Print", "Annual qty",
           "Unit price (INR, ex-GST)", "GST %", "Lead time (days)", "Remarks"]
    ws.append(hdr)
    for l in LINES:
        ws.append([l["line_id"], l["item_code"], l["description"], l["dimensions_mm"], l["ply"],
                   l["board_spec"], l["print"], l["annual_qty"], None, None, None, None])
    wq = wb.create_sheet("Questionnaire")
    wq.append(["Q#", "Question", "Your answer"])
    for q in QUESTIONNAIRE:
        wq.append([q["qid"], q["text"], None])
    wb.save(out / "quote_template.xlsx")
    return rfx


# --------------------------------------------------------------------------
# 5. Vendor 1 - Pacifica: their own Excel, USD, per 100 pcs (per 1000 for accessories)
# --------------------------------------------------------------------------

def build_v1(folder):
    T("V1")["terms"] = {"currency": "USD", "gst": "5% extra", "freight": "included (delivered Chakan)",
                         "payment_terms": "30 days", "validity": "60 days", "discount": None,
                         "lead_time_days": 18}
    wb = Workbook()
    ws = wb.active
    ws.title = "Offer"
    thin = Side(style="thin", color="999999")
    ws["A1"] = "PACIFICA PACKAGING INDIA PVT. LTD."
    ws["A1"].font = Font(bold=True, size=16, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="0B4F6C")
    ws.merge_cells("A1:J1")
    ws["A2"] = "A Pacifica Group company - Singapore | Ranjangaon MIDC, Pune"
    ws.merge_cells("A2:J2")
    ws["A4"] = "Commercial Offer No. PPI/Q/26/3381    Date: 06-Oct-2026"
    ws["A5"] = "Customer: Meghna Foods - Chakan    Attn: Ms. Ritika Deshmukh"
    ws["A6"] = "All prices in US Dollars as per Group pricing policy. See 'Terms' tab."
    ws["A6"].font = Font(italic=True, color="C00000")
    hdr = ["S.No", "PPI Code", "Product Description", "Int. Dim. LxWxH (mm)", "Ply", "Board",
           "MOQ", "UOM", "Price (USD)", "Remarks"]
    ws.append([])
    ws.append(hdr)
    hrow = ws.max_row
    for c in ws[hrow]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2E86AB")
        c.alignment = Alignment(wrap_text=True)
    # Pacifica groups by ply, largest first - not the RFx order
    order = sorted(LINES, key=lambda l: (-l["ply"], -(l["L"] * l["W"] * max(l["H"], 1))))
    sno = 1
    for l in order:
        lid = l["line_id"]
        inr = vendor_price("V1", lid)
        if l["kind"] in ("RSC",):
            unit, per = "Per 100 Nos", 100
            desc = f"RSC Carton {l['ply']} Ply, {l['print']}"
            dim = f"{l['L']}x{l['W']}x{l['H']}"
        elif l["kind"] == "TRAY":
            unit, per = "Per 100 Nos", 100
            desc = f"Die-cut Tray {l['ply']} Ply, {l['print']}"
            dim = f"{l['L']}x{l['W']}x{l['H']}"
        elif l["kind"] == "PAD":
            unit, per = "Per 1000 Nos", 1000
            desc = "Layer Pad / Separator 3 Ply"
            dim = f"{l['L']}x{l['W']}"
        else:
            unit, per = "Per 1000 Sets", 1000
            cells = 12 if l["kind"] == "PARTITION12" else 24
            desc = f"Partition {cells} Cell, 3 Ply (for your {'L11' if cells == 12 else 'L16'})"
            dim = "-"
        usd = round(inr * per / FX_USD_INR, 2)
        remarks = ""
        if lid in ("L07", "L19", "L26"):
            remarks = f"Your ref {lid}"
        ws.append([sno, f"PPI-{l['ply']}P-{dim.replace('x', '')}", desc, dim, l["ply"],
                   "Std. Kraft" if l["ply"] == 3 else ("Hi-BF Kraft" if l["ply"] == 5 else "Heavy Duty"),
                   "5,000" if per == 100 else "10,000", unit, usd, remarks])
        record("V1", lid, "quoted", usd, unit, "USD", usd / per * FX_USD_INR)
        sno += 1
        if lid == "L12":  # an alternate, cheaper option on one line
            alt = round(usd * 0.93, 2)
            ws.append([f"{sno - 1}A", f"PPI-5P-{dim.replace('x', '')}-E", desc + " - OPTION B: economy 150 GSM top liner",
                       dim, 5, "Economy Kraft", "5,000", unit, alt, "Alternate offer (spec deviation)"])
            T("V1")["lines"][lid]["alternate"] = {"native_price": alt, "note": "Option B, economy 150 GSM top liner"}
    for row in ws.iter_rows(min_row=hrow, max_row=ws.max_row):
        for c in row:
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
    for col, wdt in zip("ABCDEFGHIJ", [6, 18, 44, 18, 5, 13, 8, 14, 11, 30]):
        ws.column_dimensions[col].width = wdt
    ws.append([])
    ws.append(["", "", "Note: Prices are exclusive of GST @ 5%. Freight to Chakan included."])

    t = wb.create_sheet("Terms")
    for r in [["Commercial Terms"], [], ["Currency", "USD (invoice in INR at SBI TT selling rate on invoice date)"],
              ["Price basis", "Delivered Chakan, freight included"], ["GST", "5% extra as applicable"],
              ["Payment", "30 days from invoice"], ["Validity", "60 days from offer date"],
              ["Lead time", "18 days from PO"], ["Price variation", "Quarterly revision linked to kraft paper index (+/- 5% band)"]]:
        t.append(r)
    t.column_dimensions["A"].width = 18
    t.column_dimensions["B"].width = 70

    q = wb.create_sheet("Vendor Questionnaire")
    ans = {"Q1": "Yes - certificate attached (valid till Aug 2028)", "Q2": "Yes, in-house lab (BCT, ECT, Bursting, Cobb)",
           "Q3": "900", "Q4": "0.6", "Q5": "18", "Q6": "Yes", "Q7": "65", "Q8": "40+", "Q9": "38",
           "Q10": "No - 30 days standard for Group; 45 days negotiable"}
    q.append(["Question", "Response"])
    for qq in QUESTIONNAIRE:
        q.append([qq["text"], ans[qq["qid"]]])
    q.column_dimensions["A"].width = 80
    q.column_dimensions["B"].width = 55
    wb.save(folder / "Pacifica_Offer_PPI-Q-26-3381.xlsx")
    T("V1")["questionnaire"] = {"Q1": "yes", "Q2": "yes", "Q3": 900, "Q4": 0.6, "Q5": 18, "Q6": "yes",
                                 "Q7": 65, "Q8": 40, "Q9": 38, "Q10": "no", "iso_valid_until": "2028-08-14",
                                 "knockout_pass": True}
    iso_certificate(folder / "Pacifica_ISO9001.pdf", VENDORS["V1"]["name"], "QCS-QM-21877",
                    "15-Aug-2025", "14-Aug-2028", "Manufacture of corrugated boxes and partitions")
    write_email(folder, "V1", f"RE: {RFX_ID} - Pacifica commercial offer", "2026-10-06T17:42:00+05:30",
                "Dear Ms. Deshmukh,\n\nPlease find attached our commercial offer against your RFQ. "
                "As discussed, we have used our standard Group format; all items are covered including "
                "an economy option for L12.\n\nQuestionnaire responses are in the third tab and our ISO "
                "certificate is attached.\n\nBest regards,\nDaniel Tan\nRegional Sales, Pacifica Packaging India",
                ["Pacifica_Offer_PPI-Q-26-3381.xlsx", "Pacifica_ISO9001.pdf"])


# --------------------------------------------------------------------------
# 6. Vendor 2 - Shree Ganesh: letterhead PDF, Rs/pc, 27 of 30 lines,
#    discount hidden in a footnote, freight 2% extra
# --------------------------------------------------------------------------

def build_v2(folder):
    skipped = {"L26", "L29", "L30"}
    T("V2")["terms"] = {"currency": "INR", "gst": "5% extra", "freight": "2% extra on basic value",
                         "payment_terms": "45 days", "validity": "90 days",
                         "discount": "4% on total invoice value if annual business exceeds Rs. 50 lakh",
                         "lead_time_days": 15}
    path = folder / "SGPI_Quotation_SG-2627-118.pdf"
    c = canvas.Canvas(str(path), pagesize=A4)
    w, h = A4

    def letterhead():
        c.setFillColor(colors.HexColor("#8B1E1E"))
        c.rect(0, h - 32 * mm, w, 32 * mm, fill=1, stroke=0)
        c.setFillColor(colors.white)
        c.setFont("Times-Bold", 20)
        c.drawString(18 * mm, h - 15 * mm, "SHREE GANESH PACKAGING INDUSTRIES")
        c.setFont("Times-Roman", 9.5)
        c.drawString(18 * mm, h - 22 * mm, "Manufacturers of Corrugated Boxes & Sheets  |  Plot J-221, MIDC Bhosari, Pune 411026")
        c.drawString(18 * mm, h - 27 * mm, "GSTIN 27AAXFS0000X1Z9 (sample)  |  sales@shreeganeshpack.example  |  +91 20 0000 0000")
        c.setFillColor(colors.black)

    letterhead()
    c.setFont("Times-Bold", 12)
    c.drawString(18 * mm, h - 42 * mm, "QUOTATION")
    c.setFont("Times-Roman", 10)
    c.drawString(18 * mm, h - 49 * mm, "Ref: SG/2627/118")
    c.drawRightString(w - 18 * mm, h - 49 * mm, "Date: 08/10/2026")
    c.drawString(18 * mm, h - 56 * mm, f"To: {BUYER['name']}, Chakan  -  Kind Attn: Ms. Ritika Deshmukh")
    c.drawString(18 * mm, h - 62 * mm, f"Sub: Your enquiry {RFX_ID} for corrugated boxes")
    cols = [18, 30, 52, 112, 147, 158, 182]
    y = h - 72 * mm
    c.setFont("Times-Bold", 9)
    for x, t in zip(cols, ["Sr", "Item Code", "Description", "Size (mm)", "Ply", "Rate*", ""]):
        c.drawString(x * mm, y, t)
    c.drawString(158 * mm, y - 3.5 * mm, "(Rs./pc)")
    c.line(18 * mm, y - 5 * mm, w - 18 * mm, y - 5 * mm)
    y -= 10 * mm
    c.setFont("Times-Roman", 9)
    sr = 1
    for l in LINES:
        lid = l["line_id"]
        if lid in skipped:
            record("V2", lid, "not_quoted", None, None, "INR", None, "silently omitted")
            continue
        rate = round(vendor_price("V2", lid), 2)
        record("V2", lid, "quoted", rate, "per piece", "INR", rate)
        c.drawString(cols[0] * mm, y, str(sr))
        c.drawString(cols[1] * mm, y, l["item_code"])
        c.drawString(cols[2] * mm, y, l["description"][:34])
        c.drawString(cols[3] * mm, y, l["dimensions_mm"])
        c.drawString(cols[4] * mm, y, str(l["ply"]))
        c.drawRightString(172 * mm, y, f"{rate:,.2f}")
        y -= 5.6 * mm
        sr += 1
    c.line(18 * mm, y + 2 * mm, w - 18 * mm, y + 2 * mm)
    y -= 4 * mm
    c.setFont("Times-Bold", 9.5)
    c.drawString(18 * mm, y, "Terms & Conditions:")
    c.setFont("Times-Roman", 9)
    for t in ["1. GST @ 5% extra as applicable.", "2. Freight: 2% extra on basic value, delivered Chakan.",
              "3. Payment: 45 days from date of invoice.", "4. Validity: 90 days.",
              "5. Delivery: 15 days from receipt of confirmed PO."]:
        y -= 5 * mm
        c.drawString(20 * mm, y, t)
    c.setFont("Times-Italic", 7)
    c.drawString(18 * mm, 16 * mm, "* Rates are basic per piece. A trade discount of 4% will be extended on the total invoice value where "
                                   "annual business placed exceeds Rs. 50 lakh.")
    c.drawString(18 * mm, 12 * mm, "  Rates subject to revision in case kraft paper prices change by more than 8%.  E&OE.")
    c.showPage()

    letterhead()
    c.setFont("Times-Bold", 12)
    c.drawString(18 * mm, h - 45 * mm, "Annexure - Supplier Questionnaire")
    ans = [("ISO 9001", "Yes, certificate enclosed"), ("Lot testing BCT/ECT/Burst", "Yes - in-house lab"),
           ("Capacity", "320 MT / month"), ("Rejection rate (12 months)", "1.2 %"), ("Lead time", "15 days"),
           ("Food-grade inks", "Yes (water-based)"), ("Recycled content", "70 %"), ("FMCG customers", "12"),
           ("Distance from Chakan", "24 km"), ("60-day payment", "No, 45 days")]
    y = h - 58 * mm
    c.setFont("Times-Roman", 10)
    for i, (k, v) in enumerate(ans, 1):
        c.drawString(20 * mm, y, f"{i}. {k}")
        c.drawString(90 * mm, y, v)
        y -= 7 * mm
    c.drawString(18 * mm, y - 15 * mm, "For Shree Ganesh Packaging Industries")
    c.drawString(18 * mm, y - 30 * mm, "Prakash Kulkarni (Proprietor)")
    c.save()
    T("V2")["questionnaire"] = {"Q1": "yes", "Q2": "yes", "Q3": 320, "Q4": 1.2, "Q5": 15, "Q6": "yes",
                                 "Q7": 70, "Q8": 12, "Q9": 24, "Q10": "no", "iso_valid_until": "2027-06-30",
                                 "knockout_pass": True}
    iso_certificate(folder / "SGPI_ISO_Certificate.pdf", VENDORS["V2"]["name"], "QCS-QM-19022",
                    "01-Jul-2024", "30-Jun-2027", "Manufacture and supply of corrugated boxes")
    write_email(folder, "V2", f"Quotation against {RFX_ID}", "2026-10-08T11:05:00+05:30",
                "Respected Madam,\n\nWith reference to your enquiry we are pleased to submit our best rates. "
                "Kindly find attached quotation and ISO certificate.\n\nThanking you,\nPrakash Kulkarni\n"
                "Shree Ganesh Packaging Industries",
                ["SGPI_Quotation_SG-2627-118.pdf", "SGPI_ISO_Certificate.pdf"])


# --------------------------------------------------------------------------
# 7. Vendor 3 - Konkan Kraft: Word letter, commercials in prose,
#    per-thousand for accessories, a spec deviation, "same as" references
# --------------------------------------------------------------------------

def build_v3(folder):
    T("V3")["terms"] = {"currency": "INR", "gst": "5% extra", "freight": "included (delivered Chakan)",
                         "payment_terms": "45 days", "validity": "till 31-Mar-2027", "discount": None,
                         "lead_time_days": 21}
    p = {lid: round(vendor_price("V3", lid), 2) for lid in LINE}
    p["L17"] = p["L18"] = p["L16"]
    per_k = {lid: round(p[lid] * 1000, -1) for lid in ("L27", "L28", "L29", "L30")}
    for lid in LINE:
        if lid in per_k:
            record("V3", lid, "quoted", per_k[lid], "per 1000", "INR", per_k[lid] / 1000)
        else:
            note = "quoted as 'same rate as L16'" if lid in ("L17", "L18") else ""
            if lid in ("L05", "L06", "L07", "L08"):
                note = "spec deviation: 140 GSM top liner instead of 150"
            record("V3", lid, "quoted", p[lid], "per box", "INR", p[lid], note)
    T("V3")["lines"]["L05"]["deviation"] = T("V3")["lines"]["L06"]["deviation"] = True
    T("V3")["lines"]["L07"]["deviation"] = T("V3")["lines"]["L08"]["deviation"] = True

    d = Document()
    st = d.styles["Normal"]
    st.font.name = "Georgia"
    st.font.size = Pt(10.5)
    d.add_paragraph("KONKAN KRAFT BOXES LLP").runs[0].bold = True
    d.add_paragraph("Survey No. 77, Talegaon Dabhade Industrial Area, Pune 410507")
    d.add_paragraph("Date: 9th October 2026")
    d.add_paragraph(f"To,\nMs. Ritika Deshmukh\nCategory Manager - Packaging\n{BUYER['name']}, Chakan")
    d.add_paragraph(f"Subject: Our offer against your RFQ {RFX_ID}")
    d.add_paragraph("Dear Ms. Deshmukh,")
    d.add_paragraph(
        "Thank you for considering Konkan Kraft for your FY 2026-27 requirement. We have studied the "
        "specifications carefully and are happy to offer our most competitive rates, all of which are on a "
        "delivered-Chakan basis (freight is included in the rates below) and exclusive of GST, which will be "
        "charged at 5%.")
    d.add_paragraph(
        f"For the 3-ply shipper range, our rates per box are: L01 at ₹{p['L01']:.2f}, L02 at ₹{p['L02']:.2f}, "
        f"L03 at ₹{p['L03']:.2f} and L04 at ₹{p['L04']:.2f}. For L05 through L08 we would like to propose a "
        f"140 GSM top liner in place of the specified 150 GSM, which in our experience gives identical BCT "
        f"performance for these sizes; on that basis L05 is ₹{p['L05']:.2f}, L06 ₹{p['L06']:.2f}, "
        f"L07 ₹{p['L07']:.2f} and L08 ₹{p['L08']:.2f}. The two larger 3-ply sizes, L09 and L10, come to "
        f"₹{p['L09']:.2f} and ₹{p['L10']:.2f} respectively.")
    d.add_paragraph(
        f"Moving to the 5-ply master cartons: L11 ₹{p['L11']:.2f}, L12 ₹{p['L12']:.2f}, L13 ₹{p['L13']:.2f}, "
        f"L14 ₹{p['L14']:.2f} and L15 ₹{p['L15']:.2f}. L16 is ₹{p['L16']:.2f}, and since L17 and L18 run on "
        f"the same deckle we can hold them at the same rate as L16. The gift shipper L19 with 4-colour print is "
        f"₹{p['L19']:.2f}; L20 ₹{p['L20']:.2f}, L21 ₹{p['L21']:.2f} and L22 ₹{p['L22']:.2f}.")
    d.add_paragraph(
        f"For the 7-ply heavy-duty range we offer L23 at ₹{p['L23']:.2f}, L24 at ₹{p['L24']:.2f}, "
        f"L25 at ₹{p['L25']:.2f} and the new export carton L26 at ₹{p['L26']:.2f} per box.")
    d.add_paragraph(
        f"Accessories are priced per thousand pieces: the 12-cell partition (L27) at ₹{per_k['L27']:,.0f} per "
        f"thousand sets, the 24-cell partition (L28) at ₹{per_k['L28']:,.0f}, layer pads (L29) at "
        f"₹{per_k['L29']:,.0f} and the die-cut tray (L30) at ₹{per_k['L30']:,.0f} per thousand.")
    d.add_paragraph(
        "Our payment expectation is 45 days from invoice. Rates are valid until 31st March 2027, after which "
        "we would request a review linked to kraft paper prices. Standard lead time is 21 days.")
    d.add_paragraph(
        "On your questionnaire: yes, we are ISO 9001 certified (copy enclosed). Every lot is tested for "
        "BCT and bursting strength - we do this through an NABL-accredited lab in Chakan as we are setting up "
        "our own lab next year. Our installed capacity is approximately 150-200 MT per month depending on the "
        "product mix, and our customer rejection rate over the last twelve months has been 1.4%. Inks are "
        "water-based and IS 15495 compliant. Recycled content is around 60%, we serve 9 food and FMCG "
        "customers, and our plant is about 31 km from Chakan. We are open to discussing 60-day terms for an "
        "annual contract of meaningful size.")
    d.add_paragraph("We look forward to a long association.\n\nWarm regards,\nAnjali Sawant\nCommercial Head, Konkan Kraft Boxes LLP")
    d.save(folder / "KonkanKraft_Offer_Letter.docx")
    T("V3")["questionnaire"] = {"Q1": "yes", "Q2": "yes (accredited third-party lab)", "Q3": "150-200 (range)",
                                 "Q4": 1.4, "Q5": 21, "Q6": "yes", "Q7": 60, "Q8": 9, "Q9": 31,
                                 "Q10": "negotiable", "iso_valid_until": "2027-11-30", "knockout_pass": True,
                                 "note": "Q3 is a range whose lower bound equals the threshold"}
    iso_certificate(folder / "KonkanKraft_ISO9001.pdf", VENDORS["V3"]["name"], "QCS-QM-20456",
                    "01-Dec-2024", "30-Nov-2027", "Design and manufacture of corrugated packaging")
    write_email(folder, "V3", f"Konkan Kraft - offer for {RFX_ID}", "2026-10-09T19:20:00+05:30",
                "Dear Ritika,\n\nPlease find our offer letter attached, along with our ISO certificate. Happy to "
                "come over to Chakan to walk you through it.\n\nRegards,\nAnjali",
                ["KonkanKraft_Offer_Letter.docx", "KonkanKraft_ISO9001.pdf"])


# --------------------------------------------------------------------------
# 8. Vendor 4 - Pioneer: phone photo of a printed rate card (by size, not by
#    line id), GST-inclusive, 28 lines + 2 sizes we never asked for,
#    expired ISO certificate
# --------------------------------------------------------------------------

def perspective_coeffs(src, dst):
    a = []
    for (x, y), (u, v) in zip(dst, src):
        a.append([x, y, 1, 0, 0, 0, -u * x, -u * y])
        a.append([0, 0, 0, x, y, 1, -v * x, -v * y])
    b = np.array(src).reshape(8)
    return np.linalg.solve(np.array(a, dtype=float), b).tolist()


def build_v4(folder):
    missing = {"L19", "L30"}
    T("V4")["terms"] = {"currency": "INR", "gst": "inclusive (5%)", "freight": "included up to 50 km",
                         "payment_terms": "advance 30%, balance 15 days", "validity": "not stated",
                         "discount": None, "lead_time_days": 10}
    rows = []
    for l in LINES:
        lid = l["line_id"]
        if lid in missing:
            record("V4", lid, "not_quoted", None, None, "INR", None)
            continue
        ex = vendor_price("V4", lid)
        if l["kind"] == "PAD":
            incl = round(ex * 1.05 * 100, 0)
            record("V4", lid, "quoted", incl, "per 100 (incl. GST)", "INR", incl / 100 / 1.05)
            rows.append((f"Pad {l['L']}x{l['W']}", "3", f"{incl:,.0f} /100"))
        elif l["kind"].startswith("PARTITION"):
            incl = round(ex * 1.05, 1)
            cells = 12 if l["kind"] == "PARTITION12" else 24
            record("V4", lid, "quoted", incl, "per set (incl. GST)", "INR", incl / 1.05)
            rows.append((f"Partition {cells} cell", "3", f"{incl:.1f}"))
        else:
            incl = round(ex * 1.05, 1)
            record("V4", lid, "quoted", incl, "per box (incl. GST)", "INR", incl / 1.05)
            rows.append((l["dimensions_mm"].replace(" ", ""), str(l["ply"]), f"{incl:.1f}"))
    rows += [("300x300x300", "5", "21.5"), ("610x410x300", "5", "44.0")]  # never asked for
    rows.sort(key=lambda r: (r[1], r[0]))

    W, H = 1300, 280 + 52 * len(rows) + 260
    card = Image.new("RGB", (W, H), (250, 247, 238))
    dr = ImageDraw.Draw(card)
    dr.text((W // 2, 60), "PIONEER PACKWELL", font=font("Arial Black.ttf", 64), fill=(20, 40, 110), anchor="mm")
    dr.text((W // 2, 125), "Corrugated Boxes  *  Bhosari, Pune  *  Ph. 98XXX XXXXX", font=font("Arial.ttf", 28), fill=(40, 40, 40), anchor="mm")
    dr.text((W // 2, 185), "RATE CARD  2026-27", font=font("Arial Bold.ttf", 40), fill=(170, 20, 20), anchor="mm")
    y = 240
    dr.rectangle([60, y, W - 60, y + 52], fill=(20, 40, 110))
    for x, t in [(90, "SIZE (mm)"), (640, "PLY"), (880, "RATE Rs. (incl. GST)")]:
        dr.text((x, y + 12), t, font=font("Arial Bold.ttf", 28), fill="white")
    y += 52
    f = font("Courier New Bold.ttf", 32)
    for i, (sz, ply, rate) in enumerate(rows):
        if i % 2:
            dr.rectangle([60, y, W - 60, y + 52], fill=(236, 232, 220))
        dr.text((90, y + 10), sz, font=f, fill=(15, 15, 15))
        dr.text((655, y + 10), ply, font=f, fill=(15, 15, 15))
        dr.text((900, y + 10), rate, font=f, fill=(15, 15, 15))
        y += 52
    dr.line([60, y, W - 60, y], fill=(20, 40, 110), width=3)
    dr.text((80, y + 30), "* Rates incl. GST. Transport free up to 50 km.", font=font("Arial.ttf", 26), fill=(40, 40, 40))
    dr.text((80, y + 70), "* 30% advance with PO, balance 15 days.", font=font("Arial.ttf", 26), fill=(40, 40, 40))
    dr.text((80, y + 110), "* Pads rate per 100 nos.", font=font("Arial.ttf", 26), fill=(40, 40, 40))
    (DATA / "print").mkdir(exist_ok=True)
    card.save(DATA / "print" / "pioneer_rate_card_flat.png")  # print it and photograph it for a real photo

    # simulate a phone photo: dark desk, angled sheet, uneven light, blur, jpeg
    bg = Image.new("RGB", (int(W * 1.45), int(H * 1.25)), (62, 48, 38))
    noise = Image.effect_noise(bg.size, 18).convert("RGB")
    bg = Image.blend(bg, noise, 0.12)
    BW, BH = bg.size
    dst = [(BW * 0.20, BH * 0.07), (BW * 0.86, BH * 0.13), (BW * 0.80, BH * 0.95), (BW * 0.10, BH * 0.88)]
    src = [(0, 0), (W, 0), (W, H), (0, H)]
    warped = card.transform(bg.size, Image.PERSPECTIVE, perspective_coeffs(src, dst), Image.BICUBIC, fillcolor=(0, 0, 0))
    mask = Image.new("L", (W, H), 255).transform(bg.size, Image.PERSPECTIVE, perspective_coeffs(src, dst), Image.BICUBIC, fillcolor=0)
    bg.paste(warped, (0, 0), mask)
    grad = Image.linear_gradient("L").rotate(-35, expand=True).resize(bg.size)
    shade = Image.new("RGB", bg.size, (0, 0, 0))
    bg = Image.composite(bg, shade, grad.point(lambda v: 150 + v * 105 // 255))
    bg = bg.rotate(-3, resample=Image.BICUBIC, fillcolor=(40, 30, 25)).filter(ImageFilter.GaussianBlur(1.1))
    bg = bg.resize((int(BW * 0.62), int(BH * 0.62)))
    bg.save(folder / "IMG_20261007_164512.jpg", quality=72)

    T("V4")["questionnaire"] = {"Q1": "yes (but attached certificate expired 31-Mar-2025)", "Q2": "yes",
                                 "Q3": 180, "Q4": "below 2% (not a number)", "Q5": 10, "Q6": "not answered",
                                 "Q7": "not answered", "Q8": "many", "Q9": 22, "Q10": "no",
                                 "iso_valid_until": "2025-03-31", "knockout_pass": False,
                                 "note": "Says ISO yes, but the attached certificate has expired; Q4 does not meet <=1.5%"}
    iso_certificate(folder / "ISO_cert.pdf", VENDORS["V4"]["name"], "QCS-QM-15530",
                    "01-Apr-2022", "31-Mar-2025", "Manufacture of corrugated boxes")
    write_email(folder, "V4", "rates", "2026-10-07T16:51:00+05:30",
                "madam pls find rate card photo. all sizes available. iso yes attached. testing yes we check "
                "every lot. capacity 180 ton month. rejection below 2%. delivery 10 days. payment 30% advance.\n"
                "regards salim\nPioneer Packwell\nSent from my phone",
                ["IMG_20261007_164512.jpg", "ISO_cert.pdf"])


# --------------------------------------------------------------------------
# 9. Vendor 5 - Maruti (incumbent): a one-line email
# --------------------------------------------------------------------------

def build_v5(folder):
    T("V5")["terms"] = {"currency": "INR", "gst": "not stated", "freight": "extra (amount not stated)",
                         "payment_terms": "not stated", "validity": "not stated", "discount": None,
                         "lead_time_days": None}
    for l in LINES:
        lid = l["line_id"]
        if l["ply"] == 3 and l["kind"] == "RSC":
            record("V5", lid, "quoted_per_kg", 38.0, "per kg", "INR", 38.0 * l["nominal_weight_kg"],
                   "per-kg rate x RFx nominal weight")
        elif l["ply"] == 5:
            record("V5", lid, "quoted_per_kg", 42.0, "per kg", "INR", 42.0 * l["nominal_weight_kg"],
                   "per-kg rate x RFx nominal weight")
        elif lid in LAST_YEAR:
            record("V5", lid, "same_as_last_year", LAST_YEAR[lid], "per piece", "INR", LAST_YEAR[lid],
                   "resolved from FY25-26 PO history")
        else:
            record("V5", lid, "unresolved", None, None, "INR", None,
                   "'same as last year' but this SKU was never bought before")
    T("V5")["questionnaire"] = {"knockout_pass": False, "note": "No questionnaire response at all"}
    write_email(folder, "V5", f"Re: {RFX_ID}", "2026-10-10T22:14:00+05:30",
                "Ritika ji,\n\n₹42/kg for the 5-ply, 38 for the 3-ply, rest same as last year, freight extra.\n\n"
                "Rajesh\n\nSent from my iPhone", [])


# --------------------------------------------------------------------------
# 10. Last year's purchase orders
# --------------------------------------------------------------------------

def build_history():
    out = fresh(DATA / "history")
    rows = []
    po = 4500120
    for q, month in enumerate(["2025-05", "2025-08", "2025-11", "2026-02"]):
        for l in LINES:
            lid = l["line_id"]
            if lid not in LAST_YEAR:
                continue
            vendor = VENDORS["V5"]["name"]
            price = LAST_YEAR[lid]
            if lid in ("L11", "L12", "L13") and q % 2 == 1:  # second source on a few lines
                vendor = VENDORS["V2"]["name"]
                price = round(price * 1.02, 2)
            rows.append({"po_number": po, "po_date": f"{month}-{rng.randint(3, 26):02d}", "vendor": vendor,
                         "item_code": l["item_code"], "description": l["description"],
                         "qty": int(l["annual_qty"] / 4 * rng.uniform(0.85, 1.15) // 100 * 100),
                         "uom": l["uom"], "unit_price_inr": price, "gst_pct": 12 if month < "2025-09" else 5})
            po += 1
    with open(out / "po_history_FY2025-26.csv", "w", newline="") as fh:
        wr = csv.DictWriter(fh, fieldnames=list(rows[0]))
        wr.writeheader()
        wr.writerows(rows)
    truth["last_year_price"] = LAST_YEAR


# --------------------------------------------------------------------------

def main():
    rfx = build_rfx()
    inbox = fresh(DATA / "inbox")
    for vid, fn in [("V1", build_v1), ("V2", build_v2), ("V3", build_v3), ("V4", build_v4), ("V5", build_v5)]:
        folder = inbox / f"{vid}_{VENDORS[vid]['short'].replace(' ', '')}"
        folder.mkdir()
        fn(folder)
    build_history()
    ev = fresh(DATA / "eval")
    truth["generated_on"] = date.today().isoformat()
    (ev / "ground_truth.json").write_text(json.dumps(truth, indent=2, ensure_ascii=False))

    # sanity summary
    value = sum(l["annual_qty"] * LAST_YEAR.get(l["line_id"], l["_base"]) for l in LINES)
    print(f"RFx lines: {len(rfx['lines'])}  |  est. annual value at last-year prices: Rs {value / 1e7:.2f} crore")
    for vid, v in truth["vendors"].items():
        st = {}
        for x in v["lines"].values():
            st[x["status"]] = st.get(x["status"], 0) + 1
        tot = sum(LINE[k]["annual_qty"] * x["inr_per_unit_ex_gst"] for k, x in v["lines"].items() if x["inr_per_unit_ex_gst"])
        print(f"  {vid} {VENDORS[vid]['short']:13} {st}  value(quoted lines) Rs {tot / 1e7:.2f} cr")


if __name__ == "__main__":
    main()
