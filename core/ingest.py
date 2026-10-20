"""Turn whatever a vendor sent into (a) parts the model can read and
(b) plain text we can search later to verify the model's citations.

Spreadsheets and Word files are rendered to text with cell / paragraph
addresses, so the model can cite "Offer!R12" or "para 8" and we can check it.
PDFs and photos go to the model natively (it reads layout and pixels);
PDF text is also extracted for verification.
"""
import json
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from google.genai import types
from openpyxl import load_workbook
from pypdf import PdfReader

from . import config

MIME = {".pdf": "application/pdf", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png",
        ".webp": "image/webp", ".heic": "image/heic"}


@dataclass
class SourceDoc:
    file: str
    kind: str  # email | xlsx | docx | pdf | image | text
    text: str = ""  # searchable text (empty for images)
    path: Path | None = None
    pages: list = field(default_factory=list)  # pdf page texts / docx paragraphs

    def parts(self):
        header = f"\n===== FILE: {self.file} ({self.kind}) =====\n"
        if self.kind in ("pdf", "image"):
            return [types.Part.from_text(text=header),
                    types.Part.from_bytes(data=self.path.read_bytes(), mime_type=MIME[self.path.suffix.lower()])]
        return [types.Part.from_text(text=header + self.text)]


def _xlsx(path):
    wb = load_workbook(path, data_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"--- Sheet: {ws.title} ---")
        for row in ws.iter_rows():
            cells = [f"{c.column_letter}={c.value}" for c in row if c.value not in (None, "")]
            if cells:
                out.append(f"{ws.title}!R{row[0].row}: " + " | ".join(cells))
    return "\n".join(out)


def _docx(path):
    d = Document(path)
    paras = [p.text for p in d.paragraphs if p.text.strip()]
    lines = [f"[para {i}] {t}" for i, t in enumerate(paras, 1)]
    for ti, t in enumerate(d.tables, 1):
        for ri, r in enumerate(t.rows, 1):
            lines.append(f"[table {ti} row {ri}] " + " | ".join(c.text for c in r.cells))
    return "\n".join(lines), paras


def _pdf(path):
    pages = [(p.extract_text() or "") for p in PdfReader(path).pages]
    return "\n".join(f"--- page {i} ---\n{t}" for i, t in enumerate(pages, 1)), pages


def load_vendor(folder: Path):
    meta = json.loads((folder / "email.json").read_text())
    body = (f"From: {meta['from_name']} <{meta['from_email']}> ({meta['from_company']})\n"
            f"Received: {meta['received_at']}\nSubject: {meta['subject']}\n\n{meta['body']}")
    docs = [SourceDoc(file="email body", kind="email", text=body)]
    for name in meta["attachments"]:
        p = folder / name
        ext = p.suffix.lower()
        if ext in (".xlsx", ".xlsm"):
            docs.append(SourceDoc(name, "xlsx", _xlsx(p), p))
        elif ext == ".docx":
            text, paras = _docx(p)
            docs.append(SourceDoc(name, "docx", text, p, paras))
        elif ext == ".pdf":
            text, pages = _pdf(p)
            docs.append(SourceDoc(name, "pdf", text, p, pages))
        elif ext in MIME:
            docs.append(SourceDoc(name, "image", "", p))
        else:
            docs.append(SourceDoc(name, "text", p.read_text(errors="ignore"), p))
    return meta, docs


def vendor_folders():
    return sorted(p for p in config.INBOX.iterdir() if p.is_dir() and (p / "email.json").exists())


def vendor_id(folder: Path):
    return folder.name.split("_")[0]


# ---- adding responses from the UI ---------------------------------------------
UPLOAD_TYPES = ["xlsx", "xlsm", "pdf", "docx", "jpg", "jpeg", "png", "webp", "heic", "txt", "csv", "eml"]
USER_MARK = ".user_changed"  # vendors added or edited in the UI may be re-read even in demo mode


def _safe_name(name):
    keep = "".join(c if c.isalnum() or c in "._- " else "_" for c in Path(name).name).strip()
    return keep or "file"


def add_vendor(company, contact, email, subject, body, files, received_at):
    """files: list of (filename, bytes). Returns the new folder."""
    import re
    nums = [int(m.group(1)) for p in config.INBOX.iterdir() if (m := re.match(r"V(\d+)_", p.name))]
    vid = f"V{max(nums, default=0) + 1}"
    short = re.sub(r"[^A-Za-z0-9]", "", company.split()[0]) or "Vendor"
    folder = config.INBOX / f"{vid}_{short}"
    folder.mkdir(parents=True)
    names = []
    for fname, data in files:
        n = _safe_name(fname)
        (folder / n).write_bytes(data)
        names.append(n)
    meta = {"from_name": contact or company, "from_company": company, "from_email": email or "",
            "to": "buyer", "subject": subject or "Quotation", "received_at": received_at,
            "body": body or "(no email text - see attachments)", "attachments": names}
    (folder / "email.json").write_text(json.dumps(meta, indent=2, ensure_ascii=False))
    (folder / USER_MARK).write_text("added")
    return folder


def add_attachment(folder: Path, fname, data):
    n = _safe_name(fname)
    (folder / n).write_bytes(data)
    meta = json.loads((folder / "email.json").read_text())
    if n not in meta["attachments"]:
        meta["attachments"].append(n)
    (folder / "email.json").write_text(json.dumps(meta, indent=2, ensure_ascii=False))
    (folder / USER_MARK).write_text("edited")
    return n


def user_changed(folder: Path):
    return (folder / USER_MARK).exists()


def delete_vendor(folder: Path):
    import shutil
    vid = vendor_id(folder)
    shutil.rmtree(folder)
    ex = config.EXTRACTIONS / f"{vid}.json"
    if ex.exists():
        ex.unlink()
